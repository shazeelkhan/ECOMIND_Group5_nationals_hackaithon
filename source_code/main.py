import sys
import os
import sqlite3
import cv2
import numpy as np
from sklearn.ensemble import IsolationForest, RandomForestRegressor, RandomForestClassifier
import matplotlib
matplotlib.use('QtAgg')
from matplotlib.figure import Figure
from matplotlib.backends.backend_qtagg import FigureCanvasQTAgg as FigureCanvas
from datetime import datetime

# Document export
from docx import Document   # type: ignore
from docx.shared import Pt, RGBColor, Inches # type: ignore
from docx.enum.text import WD_ALIGN_PARAGRAPH # type: ignore

from PyQt6.QtCore import Qt, QTimer
from PyQt6.QtGui import QFont, QPixmap, QColor, QPainter, QPen, QBrush, QRadialGradient
from PyQt6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QLabel, QPushButton, QStackedWidget, QFrame, QGridLayout,
    QLineEdit, QComboBox, QDoubleSpinBox, QCheckBox, QFileDialog,
    QMessageBox, QTableWidget, QTableWidgetItem, QProgressBar,
    QTextEdit, QDialog, QScrollArea, QMenu, QSizePolicy, QSlider
)

# ============================================================
# CONFIGURATION
# ============================================================
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_NAME = os.path.join(BASE_DIR, "ecomind.db")
IMAGE_FOLDER = os.path.join(BASE_DIR, "waste_images")
APP_NAME = "EcoMind AI"

CATEGORIES = [
    "Food",
    "Plastic",
    "Paper",
    "Glass",
    "Metal",
    "E-waste"
]

# Alert when a single record or category total exceeds this (kg)
HIGH_WASTE_THRESHOLD_KG = 10.0

# Predefined suggestions when waste goes high
HIGH_WASTE_SUGGESTIONS = {
    "Food": [
        "Plan weekly meals to avoid over-buying and food spoilage.",
        "Compost leftover organic waste instead of landfilling.",
        "Store perishables correctly and use leftovers creatively.",
    ],
    "Plastic": [
        "Switch to reusable bags, bottles, and containers.",
        "Avoid single-use plastics; choose bulk or package-free options.",
        "Separate clean plastics for local recycling collection.",
    ],
    "Paper": [
        "Go digital for bills, notes, and documents where possible.",
        "Reuse paper for drafts before recycling.",
        "Buy recycled paper products to close the loop.",
    ],
    "Glass": [
        "Reuse glass jars for storage before recycling.",
        "Rinse glass containers so recyclers accept them.",
        "Avoid breaking glass; keep it separate from other waste.",
    ],
    "Metal": [
        "Collect cans and scrap metal for dedicated recycling.",
        "Donate usable metal items instead of discarding.",
        "Crush cans to save space and improve collection efficiency.",
    ],
    "E-waste": [
        "Take electronics to certified e-waste collection centres.",
        "Repair or refurbish devices before replacing them.",
        "Never dispose of batteries or gadgets in regular trash.",
    ],
    "General": [
        "Track waste daily and set a reduction target for next week.",
        "Segregate at source: wet, dry, and hazardous streams.",
        "Share high-waste alerts with household members for action.",
    ],
}

# ============================================================
# DATABASE
# ============================================================
class Database:
    def __init__(self):
        self.connection = sqlite3.connect(DB_NAME)
        self.create_tables()
        self.seed_default_houses()

    def create_tables(self):
        cursor = self.connection.cursor()
        # Houses table
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS houses (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL UNIQUE
            )
        """)
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS waste_records (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                record_date TEXT NOT NULL,
                category TEXT NOT NULL,
                weight REAL NOT NULL,
                recycled INTEGER NOT NULL,
                description TEXT,
                image_path TEXT,
                ai_category TEXT,
                ai_confidence REAL,
                house_id INTEGER DEFAULT 1
            )
        """)
        # Check existing columns
        cursor.execute("PRAGMA table_info(waste_records)")
        columns = {row[1] for row in cursor.fetchall()}
        
        # Add columns missing from older databases
        missing_columns = {
            "description": "TEXT",
            "image_path": "TEXT",
            "ai_category": "TEXT",
            "ai_confidence": "REAL",
            "house_id": "INTEGER DEFAULT 1"
        }
        for column, data_type in missing_columns.items():
            if column not in columns:
                cursor.execute(f"ALTER TABLE waste_records ADD COLUMN {column} {data_type}")
        self.connection.commit()

    def seed_default_houses(self):
        """Create default houses and sample data if the database is fresh."""
        cursor = self.connection.cursor()
        cursor.execute("SELECT COUNT(*) FROM houses")
        count = cursor.fetchone()[0]
        
        if count == 0:
            # 1. Seed the 5 default houses
            defaults = ["Sharma House", "Gupta House", "Patel House", "Verma House", "Ghazeel House"]
            for name in defaults:
                cursor.execute("INSERT INTO houses (name) VALUES (?)", (name,))
            self.connection.commit()
            
            # 2. Pre-load sample waste records for demonstration
            sample_records = [
                ("2026-10-01", "Food", 2.5, 1, "Daily food waste", "", "Food", 95.0, 1),
                ("2026-10-01", "Plastic", 1.2, 0, "Packaging waste", "", "Plastic", 88.0, 2),
                ("2026-10-02", "Paper", 0.8, 1, "Old newspapers", "", "Paper", 92.0, 3),
                ("2026-10-02", "Glass", 1.5, 1, "Bottles and jars", "", "Glass", 98.0, 4),
                ("2026-10-03", "Metal", 0.5, 0, "Cans", "", "Metal", 85.0, 5),
                ("2026-10-03", "E-waste", 0.2, 0, "Old cables", "", "E-waste", 70.0, 1),
                ("2026-10-04", "Food", 3.1, 1, "Leftovers", "", "Food", 90.0, 2),
                ("2026-10-04", "Plastic", 2.0, 0, "Plastic bags", "", "Plastic", 82.0, 3),
                ("2026-10-05", "Paper", 1.1, 1, "Cardboard boxes", "", "Paper", 94.0, 4),
                ("2026-10-05", "Glass", 0.9, 1, "Wine bottles", "", "Glass", 96.0, 5),
            ]
            for rec in sample_records:
                cursor.execute("""
                    INSERT INTO waste_records 
                    (record_date, category, weight, recycled, description, image_path, ai_category, ai_confidence, house_id)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, rec)
            self.connection.commit()

    def get_houses(self):
        cursor = self.connection.cursor()
        cursor.execute("SELECT id, name FROM houses ORDER BY id ASC")
        return cursor.fetchall()

    def add_house(self, name):
        name = (name or "").strip()
        if not name:
            return False, "House name cannot be empty."
        cursor = self.connection.cursor()
        try:
            cursor.execute("INSERT INTO houses (name) VALUES (?)", (name,))
            self.connection.commit()
            return True, cursor.lastrowid
        except sqlite3.IntegrityError:
            return False, "A house with this name already exists."

    def add_record(self, record_date, category, weight, recycled, description="", image_path="", ai_category="", ai_confidence=0, house_id=1):
        cursor = self.connection.cursor()
        cursor.execute("""
            INSERT INTO waste_records
            (record_date, category, weight, recycled, description, image_path, ai_category, ai_confidence, house_id)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (record_date, category, weight, int(recycled), description, image_path, ai_category, ai_confidence, int(house_id) if house_id else 1))
        self.connection.commit()

    def get_records(self, house_id=None):
        cursor = self.connection.cursor()
        query = """
            SELECT w.id, w.record_date, w.category, w.weight, w.recycled, w.description, 
                   w.image_path, w.ai_category, w.ai_confidence, w.house_id, 
                   COALESCE(h.name, 'Unknown House')
            FROM waste_records w
            LEFT JOIN houses h ON w.house_id = h.id
        """
        if house_id is not None:
            query += " WHERE w.house_id = ?"
            cursor.execute(query + " ORDER BY w.record_date DESC, w.id DESC", (house_id,))
        else:
            cursor.execute(query + " ORDER BY w.record_date DESC, w.id DESC")
        return cursor.fetchall()

    def delete_record(self, record_id):
        cursor = self.connection.cursor()
        cursor.execute("DELETE FROM waste_records WHERE id = ?", (record_id,))
        self.connection.commit()

    def clear_all(self):
        cursor = self.connection.cursor()
        cursor.execute("DELETE FROM waste_records")
        self.connection.commit()

    def close(self):
        self.connection.close()

# ============================================================
# ANALYTICS
# ============================================================
class Analytics:
    def __init__(self, database):
        self.database = database

    def records(self):
        return self.database.get_records()

    def total_waste(self):
        return sum(row[3] for row in self.records())

    def recycled_waste(self):
        return sum(row[3] for row in self.records() if row[4] == 1)

    def recycling_rate(self):
        total = self.total_waste()
        if total == 0:
            return 0
        return (self.recycled_waste() / total) * 100

    def category_breakdown(self):
        result = {}
        for row in self.records():
            category = row[2]
            weight = row[3]
            result[category] = result.get(category, 0) + weight
        return result

    def highest_category(self):
        data = self.category_breakdown()
        if not data:
            return "No data"
        return max(data, key=data.get)

    def sustainability_score(self):
        rate = self.recycling_rate()
        score = min(100, rate)
        return score

# ============================================================
# MACHINE LEARNING ENGINE
# ============================================================
class EcoMindMLEngine:
    def __init__(self, database):
        self.database = database
        self.anomaly_model = None
        self.prediction_model = None
        self.vision_model = None
        self.vision_classes = []

    def _daily_data(self):
        daily = {}
        for row in self.database.get_records():
            date = str(row[1])
            category = str(row[2])
            weight = float(row[3] or 0)
            recycled = int(row[4] or 0)
            if date not in daily:
                daily[date] = {"total": 0.0, "recycled": 0.0, "categories": {}}
            daily[date]["total"] += weight
            daily[date]["recycled"] += weight if recycled else 0.0
            daily[date]["categories"][category] = daily[date]["categories"].get(category, 0.0) + weight
        return [(d, daily[d]) for d in sorted(daily)]

    def anomaly_detection(self):
        days = self._daily_data()
        if len(days) < 5:
            return {"status": "LEARNING", "is_anomaly": False, "strength": 0.0, "message": "Need at least 5 different days of data for anomaly detection."}
        
        X = []
        for _, item in days:
            total = item["total"]
            recycle_rate = (item["recycled"] / total * 100) if total else 0
            X.append([total, recycle_rate] + [item["categories"].get(c, 0.0) for c in CATEGORIES])
        
        self.anomaly_model = IsolationForest(n_estimators=150, contamination="auto", random_state=42)
        self.anomaly_model.fit(np.asarray(X, dtype=float))
        labels = self.anomaly_model.predict(np.asarray(X, dtype=float))
        scores = self.anomaly_model.decision_function(np.asarray(X, dtype=float))
        
        last_label = int(labels[-1])
        raw = float(scores[-1])
        strength = float(np.clip(50.0 - raw * 100.0, 0.0, 100.0))
        
        if last_label == -1:
            msg = f"Unusually high/low waste pattern detected today. Anomaly strength: {strength:.0f}/100."
        else:
            msg = f"Today's waste pattern is within the learned historical range. Normality strength: {100-strength:.0f}/100."
        
        return {"status": "ACTIVE", "is_anomaly": last_label == -1, "strength": strength, "message": msg}

    def waste_prediction(self):
        days = self._daily_data()
        if len(days) < 8:
            return {"status": "LEARNING", "prediction": None, "direction": "unknown", "message": "Need at least 8 different days of data for a meaningful next-day prediction."}
        
        totals = [item["total"] for _, item in days]
        X, y = [], []
        for i in range(3, len(totals)):
            history = totals[i-3:i]
            X.append(history + [float(np.mean(history))])
            y.append(totals[i])
        
        self.prediction_model = RandomForestRegressor(n_estimators=200, random_state=42, min_samples_leaf=1)
        self.prediction_model.fit(np.asarray(X, dtype=float), np.asarray(y, dtype=float))
        
        latest = totals[-3:]
        prediction = float(self.prediction_model.predict([latest + [float(np.mean(latest))]])[0])
        prediction = max(0.0, prediction)
        latest_total = totals[-1]
        
        if prediction > latest_total * 1.08:
            direction = "rising"
        elif prediction < latest_total * 0.92:
            direction = "falling"
        else:
            direction = "stable"
            
        return {"status": "ACTIVE", "prediction": prediction, "direction": direction, "message": f"ML predicts approximately {prediction:.2f} kg of total waste for the next recorded day."}

    @staticmethod
    def _image_features(path):
        image = cv2.imread(path)
        if image is None:
            return None
        image = cv2.resize(image, (64, 64), interpolation=cv2.INTER_AREA)
        hsv = cv2.cvtColor(image, cv2.COLOR_BGR2HSV)
        gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
        features = []
        for channel in range(3):
            hist = cv2.calcHist([hsv], [channel], None, [16], [0, 256]).flatten()
            hist = hist / (hist.sum() + 1e-9)
            features.extend(hist.tolist())
        gray_hist = cv2.calcHist([gray], [0], None, [16], [0, 256]).flatten()
        gray_hist = gray_hist / (gray_hist.sum() + 1e-9)
        features.extend(gray_hist.tolist())
        mean = image.mean(axis=(0, 1)) / 255.0
        std = image.std(axis=(0, 1)) / 255.0
        features.extend(mean.tolist())
        features.extend(std.tolist())
        return np.asarray(features, dtype=np.float32)

    def train_vision(self):
        X, y = [], []
        records = self.database.get_records()
        per_class = {}
        for row in records:
            image_path = row[6]
            category = row[2]
            if not image_path or not os.path.exists(image_path):
                continue
            features = self._image_features(image_path)
            if features is None:
                continue
            X.append(features)
            y.append(category)
            per_class[category] = per_class.get(category, 0) + 1
        
        usable_classes = [c for c, n in per_class.items() if n >= 3]
        if len(usable_classes) < 2:
            self.vision_model = None
            self.vision_classes = []
            return {"status": "LEARNING", "message": "Vision model needs at least 3 saved labeled images in each of 2 categories.", "classes": per_class}
        
        filtered = [(features, label) for features, label in zip(X, y) if label in usable_classes]
        X2 = np.asarray([item[0] for item in filtered], dtype=np.float32)
        y2 = np.asarray([item[1] for item in filtered])
        
        self.vision_model = RandomForestClassifier(n_estimators=250, random_state=42, class_weight="balanced")
        self.vision_model.fit(X2, y2)
        self.vision_classes = list(self.vision_model.classes_)
        return {"status": "ACTIVE", "message": f"Vision classifier trained from {len(y2)} labeled images across {len(self.vision_classes)} categories.", "classes": per_class}

    def classify_image(self, image_path):
        result = self.train_vision()
        if self.vision_model is None:
            return "", 0.0, result["message"]
        features = self._image_features(image_path)
        if features is None:
            return "", 0.0, "Unable to read the image for ML classification."
        probabilities = self.vision_model.predict_proba([features])[0]
        index = int(np.argmax(probabilities))
        category = str(self.vision_model.classes_[index])
        confidence = float(probabilities[index] * 100.0)
        return category, confidence, "Trained local image classifier prediction."

    def analyze(self):
        anomaly = self.anomaly_detection()
        prediction = self.waste_prediction()
        vision = self.train_vision()
        return {
            "anomaly": anomaly,
            "prediction": prediction,
            "vision": vision,
            "active_models": sum([anomaly["status"] == "ACTIVE", prediction["status"] == "ACTIVE", vision["status"] == "ACTIVE"])
        }

# ============================================================
# COMPUTER VISION
# ============================================================
class WasteVision:
    def __init__(self):
        os.makedirs(IMAGE_FOLDER, exist_ok=True)
        self.no_detection_seconds = 5.0
        self.warmup_seconds = 1.0
        self.min_contour_area = 4500

    def _save_frame(self, frame):
        filename = "waste_" + datetime.now().strftime("%Y%m%d_%H%M%S") + ".jpg"
        saved_path = os.path.join(IMAGE_FOLDER, filename)
        if cv2.imwrite(saved_path, frame):
            return saved_path
        return None

    def _object_detected(self, frame, background_gray):
        gray = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
        gray = cv2.GaussianBlur(gray, (21, 21), 0)
        difference = cv2.absdiff(background_gray, gray)
        _, threshold = cv2.threshold(difference, 30, 255, cv2.THRESH_BINARY)
        threshold = cv2.dilate(threshold, None, iterations=2)
        contours, _ = cv2.findContours(threshold, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
        largest_area = 0
        for contour in contours:
            area = cv2.contourArea(contour)
            if area > largest_area:
                largest_area = area
        return largest_area >= self.min_contour_area, largest_area

    def capture_image(self):
        camera = cv2.VideoCapture(0)
        if not camera.isOpened():
            QMessageBox.warning(None, "Camera Error", "Could not access the camera.\n\nCheck that your webcam is connected.")
            return None
        
        start_time = datetime.now().timestamp()
        background_gray = None
        captured_path = None
        detected = False
        detection_streak = 0
        last_status = "Initializing camera..."
        
        try:
            while True:
                success, frame = camera.read()
                if not success:
                    break
                now = datetime.now().timestamp()
                elapsed = now - start_time
                gray = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
                gray = cv2.GaussianBlur(gray, (21, 21), 0)
                
                if background_gray is None:
                    background_gray = gray.copy()
                elif elapsed < self.warmup_seconds:
                    background_gray = cv2.addWeighted(background_gray, 0.85, gray, 0.15, 0)
                
                if elapsed >= self.warmup_seconds:
                    found, area = self._object_detected(frame, background_gray)
                    if found:
                        detection_streak += 1
                        last_status = "Possible waste detected - checking..."
                    else:
                        detection_streak = 0
                        last_status = f"Searching for waste... ({elapsed:.1f}s / {self.no_detection_seconds:.0f}s)"
                    
                    if detection_streak >= 2:
                        detected = True
                        cv2.putText(frame, "WASTE OBJECT DETECTED - CAPTURING...", (20, 40), cv2.FONT_HERSHEY_SIMPLEX, 0.75, (80, 230, 170), 2)
                        cv2.imshow("EcoMind AI - Camera | SPACE = Capture | ESC = Cancel", frame)
                        cv2.waitKey(300)
                        captured_path = self._save_frame(frame)
                        break
                
                cv2.putText(frame, last_status, (20, 35), cv2.FONT_HERSHEY_SIMPLEX, 0.65, (255, 255, 255), 2)
                cv2.putText(frame, "Place waste in front of camera | ESC = Cancel", (20, frame.shape[0] - 20), cv2.FONT_HERSHEY_SIMPLEX, 0.55, (220, 220, 220), 1)
                cv2.imshow("EcoMind AI - Camera | SPACE = Capture | ESC = Cancel", frame)
                
                key = cv2.waitKey(1) & 0xFF
                if key == 32:
                    captured_path = self._save_frame(frame)
                    break
                if key == 27:
                    break
                if elapsed >= self.no_detection_seconds:
                    break
        finally:
            camera.release()
            cv2.destroyAllWindows()
        
        if detected and captured_path:
            QMessageBox.information(None, "Waste Detected", "Waste/object detected and image captured successfully.")
            return captured_path
        if not captured_path and not detected:
            QMessageBox.information(None, "No Waste Detected", "No waste detected for 5 seconds. Camera turned off.")
        return captured_path

# ============================================================
# MAIN APPLICATION
# ============================================================
class GraphWindow(QDialog):
    def __init__(self, db, parent=None):
        super().__init__(parent)
        self.db = db
        self.setWindowTitle("EcoMind AI — Waste Graphs")
        self.resize(1050, 700)

        layout = QVBoxLayout(self)
        controls = QHBoxLayout()
        controls.addWidget(QLabel("Graph:"))
        self.graph_type = QComboBox()
        self.graph_type.addItems(["Bar Graph — Category Waste", "Line Graph — Waste Trend"])
        controls.addWidget(self.graph_type)
        refresh = QPushButton("Refresh Graph")
        refresh.clicked.connect(self.draw_graph)
        controls.addWidget(refresh)
        controls.addStretch()
        layout.addLayout(controls)

        self.figure = Figure(figsize=(10, 6), tight_layout=True)
        self.canvas = FigureCanvas(self.figure)
        layout.addWidget(self.canvas)
        self.draw_graph()

    def _rows(self):
        try:
            cur = self.db.connection.cursor()
            cur.execute("SELECT record_date, category, weight FROM waste_records ORDER BY record_date ASC, id ASC")
            return cur.fetchall()
        except Exception:
            return []

    def draw_graph(self):
        rows = self._rows()
        self.figure.clear()
        ax = self.figure.add_subplot(111)
        if not rows:
            ax.text(0.5, 0.5, "No waste data available yet.", ha="center", va="center", transform=ax.transAxes)
            ax.set_axis_off()
            self.canvas.draw()
            return
        
        graph = self.graph_type.currentText()
        if graph.startswith("Bar"):
            totals = {}
            for date, category, weight in rows:
                category = str(category or "Unknown")
                totals[category] = totals.get(category, 0.0) + float(weight or 0)
            categories = list(totals.keys())
            values = [totals[c] for c in categories]
            ax.bar(categories, values)
            ax.set_title("Waste Generated by Category")
            ax.set_xlabel("Waste Category")
            ax.set_ylabel("Weight (kg)")
            ax.tick_params(axis="x", rotation=30)
        else:
            daily = {}
            for date, category, weight in rows:
                day = str(date)
                daily[day] = daily.get(day, 0.0) + float(weight or 0)
            dates = list(daily.keys())
            values = [daily[d] for d in dates]
            ax.plot(dates, values, marker="o", linewidth=2)
            ax.set_title("Waste Generation Trend")
            ax.set_xlabel("Date")
            ax.set_ylabel("Waste (kg)")
            ax.tick_params(axis="x", rotation=30)
            ax.grid(True, alpha=0.25)
        
        self.figure.tight_layout()
        self.canvas.draw()

class EcoFlowMap(QWidget):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.breakdown = {}
        self.score = 0.0
        self.phase = 0
        self.setMinimumHeight(330)
        self.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)
        self.timer = QTimer(self)
        self.timer.timeout.connect(self._animate)
        self.timer.start(90)

    def set_data(self, breakdown, score):
        self.breakdown = dict(breakdown or {})
        self.score = float(score or 0)
        self.update()

    def _animate(self):
        self.phase = (self.phase + 1) % 120
        self.update()

    def paintEvent(self, event):
        painter = QPainter(self)
        painter.setRenderHint(QPainter.RenderHint.Antialiasing)
        w, h = self.width(), self.height()
        painter.fillRect(self.rect(), QColor('#06101E'))
        
        grid_pen = QPen(QColor(16, 44, 70, 95), 1)
        painter.setPen(grid_pen)
        horizon = int(h * 0.64)
        for y in range(horizon, h + 30, 24):
            painter.drawLine(0, y, w, y)
        for x in range(-w, w * 2, 42):
            painter.drawLine(w // 2, horizon, x, h)
        
        glow = QRadialGradient(w * 0.50, h * 0.45, min(w, h) * 0.40)
        glow.setColorAt(0.0, QColor(20, 90, 130, 48))
        glow.setColorAt(0.45, QColor(10, 55, 95, 22))
        glow.setColorAt(1.0, QColor(0, 0, 0, 0))
        painter.setBrush(QBrush(glow))
        painter.setPen(Qt.PenStyle.NoPen)
        painter.drawEllipse(int(w*0.12), int(h*0.02), int(w*0.76), int(h*0.82))
        
        cx, cy = int(w * 0.50), int(h * 0.47)
        battery_shift_px = (7.0 / 25.4) * self.logicalDpiY()
        battery_py = 0.66 - (battery_shift_px / max(h, 1))
        
        positions = {
            'Food': (0.14, 0.22, '#39E6A4', 'ORGANIC'),
            'Plastic': (0.31, 0.10, '#31D8FF', 'PLASTIC'),
            'Paper': (0.64, 0.10, '#FFC857', 'PAPER'),
            'Battery Waste': (0.24, battery_py, '#FFD447', 'BATTERY WASTE'),
            'Glass': (0.84, 0.22, '#7895FF', 'GLASS'),
            'Metal': (0.79, 0.54, '#C184FF', 'METAL'),
            'E-waste': (0.61, 0.62, '#FF6FB5', 'E-WASTE'),
            'Other': (0.14, 0.43, '#61E6D2', 'OTHER'),
        }
        
        for cat, (px, py, color, label) in positions.items():
            x, y = int(w*px), int(h*py)
            pen = QPen(QColor(color), 2)
            painter.setPen(pen)
            painter.drawLine(cx, cy, x, y)
            t = ((self.phase + list(positions).index(cat)*17) % 100) / 100
            pulse_x = cx + (x-cx)*t
            pulse_y = cy + (y-cy)*t
            painter.setBrush(QBrush(QColor(color)))
            painter.setPen(Qt.PenStyle.NoPen)
            painter.drawEllipse(int(pulse_x)-3, int(pulse_y)-3, 6, 6)
            painter.setBrush(QBrush(QColor(color + '28')))
            painter.drawEllipse(x-20, y-20, 40, 40)
            painter.setBrush(QBrush(QColor('#091B2D')))
            painter.setPen(QPen(QColor(color), 2))
            painter.drawEllipse(x-13, y-13, 26, 26)
            painter.setPen(QColor(color))
            painter.setFont(QFont('Segoe UI', 8, QFont.Weight.Bold))
            painter.drawText(x-45, y+27, 90, 14, Qt.AlignmentFlag.AlignCenter, label)
        
        r = 64
        pulse = 3 + int(4 * abs(((self.phase % 40) / 20) - 1))
        painter.setBrush(QBrush(QColor(31, 216, 180, 35)))
        painter.setPen(Qt.PenStyle.NoPen)
        painter.drawEllipse(cx-r-pulse, cy-r-pulse, 2*(r+pulse), 2*(r+pulse))
        painter.setBrush(QBrush(QColor('#0B2034')))
        painter.setPen(QPen(QColor('#35D8FF'), 2.2))
        painter.drawEllipse(cx-r, cy-r, 2*r, 2*r)
        painter.setPen(QColor('#F1FBFF'))
        painter.setFont(QFont('Segoe UI', 13, QFont.Weight.Bold))
        painter.drawText(cx-r, cy-9, 2*r, 22, Qt.AlignmentFlag.AlignCenter, 'AI CORE')
        
        painter.setPen(QColor('#6E8AA5'))
        painter.setFont(QFont('Segoe UI', 7, QFont.Weight.Bold))
        painter.drawText(12, 17, 'LIVE MATERIAL FLOW  /  LOCAL INTELLIGENCE')
        painter.setPen(QColor('#3CE4B0'))
        painter.drawText(w-90, 17, 78, 12, Qt.AlignmentFlag.AlignRight, '● ONLINE')

class EcoMindApp(QMainWindow):
    def open_graphs(self):
        try:
            self.graph_window = GraphWindow(self.database, self)
            self.graph_window.exec()
        except Exception as e:
            QMessageBox.warning(self, "Graphs", f"Could not open graphs:\n{e}")

    def __init__(self):
        super().__init__()
        self.database = Database()
        self.analytics = Analytics(self.database)
        self.ml_engine = EcoMindMLEngine(self.database)
        self.vision = WasteVision()
        self.selected_image = ""
        self.ai_category = ""
        self.ai_confidence = 0
        
        self.setWindowTitle("EcoMind AI — Intelligent Waste Management System")
        screen = QApplication.primaryScreen().availableGeometry()
        width = min(1400, max(1100, int(screen.width() * 0.92)))
        height = min(820, max(650, int(screen.height() * 0.90)))
        self.resize(width, height)
        self.setMinimumSize(1050, 620)
        
        self.build_ui()
        self.load_dashboard()
        
        self.clock_timer = QTimer(self)
        self.clock_timer.timeout.connect(self.update_clock)
        self.clock_timer.start(1000)
        self.update_clock()

    def update_clock(self):
        now = datetime.now()
        self.clock_time_label.setText(now.strftime('%I:%M:%S %p'))
        self.clock_date_label.setText(now.strftime('%d %b %Y'))

    def apply_style(self):
        self.setStyleSheet("""
            QMainWindow { background-color: #050816; }
            QWidget { color: #e7fff7; font-family: Segoe UI; }
            QFrame#Sidebar { background-color: qlineargradient(x1:0, y1:0, x2:0, y2:1, stop:0 #081A2B, stop:0.55 #071321, stop:1 #10102A); border-right: 1px solid #25405E; }
            QFrame#TopBar { background-color: qlineargradient(x1:0, y1:0, x2:1, y2:0, stop:0 #081625, stop:0.55 #0A1730, stop:1 #15102B); border-bottom: 1px solid #182C45; }
            QLabel#Title { font-size: 25px; font-weight: bold; color: #7fffd4; }
            QLabel#Subtitle { color: #8AA0B6; font-size: 12px; }
            QPushButton { background-color: #0C1B2D; border: 1px solid #244563; border-radius: 8px; padding: 10px 15px; color: #DCEBFA; }
            QPushButton:hover { background-color: #172B4A; border: 1px solid #7B6CFF; }
            QPushButton#AccountButton { background: qlineargradient(x1:0, y1:0, x2:1, y2:0, stop:0 #0b1d30, stop:0.55 #10263d, stop:1 #111a35); color: #eaf7ff; border: 1px solid #244d70; border-radius: 10px; text-align: left; padding: 8px 10px; font-size: 12px; font-weight: 600; }
            QPushButton#AccountButton:hover { border: 1px solid #39d9ff; background: #132d46; }
            QPushButton#NavButton { text-align: left; padding: 13px; border: none; border-radius: 7px; }
            QPushButton#NavButton:hover { background-color: #10243A; }
            QPushButton#ActiveNav { background-color: qlineargradient(x1:0, y1:0, x2:1, y2:0, stop:0 #183D67, stop:1 #29205B); color: #7FFFE0; border-left: 3px solid #45E5FF; }
            QFrame#Card { background-color: #0B1828; border: 1px solid #193452; border-radius: 12px; }
            QLabel#CardTitle { color: #8AA1B8; font-size: 12px; }
            QLabel#CardValue { color: #F2F8FF; font-size: 25px; font-weight: bold; }
            QLabel#SectionTitle { font-size: 20px; font-weight: bold; color: #69E8FF; }
            
            /* FIX: Proper styling for inputs and dropdowns to prevent light text on white background */
            QLineEdit, QDoubleSpinBox, QTextEdit {
                background-color: #071421;
                border: 1px solid #1C3A59;
                border-radius: 7px;
                padding: 9px;
                color: white;
            }
            QComboBox {
                background-color: #071421;
                color: white;
                border: 1px solid #1C3A59;
                border-radius: 7px;
                padding: 9px;
            }
            QComboBox QAbstractItemView {
                background-color: #071421;
                color: #e7fff7;
                selection-background-color: #172B4A;
                border: 1px solid #1C3A59;
                outline: 0;
            }
            QComboBox::drop-down {
                border: none;
                background: transparent;
                width: 20px;
            }
            
            QTableWidget { background-color: #091813; border: 1px solid #193452; gridline-color: #172E48; }
            QHeaderView::section { background-color: #10243A; color: #69E8FF; padding: 8px; border: none; }
            QProgressBar { background-color: #10231e; border-radius: 8px; height: 14px; text-align: center; }
            QProgressBar::chunk { background-color: #37D7A5; border-radius: 8px; }
        """)

    def build_ui(self):
        self.apply_style()
        root = QWidget()
        main_layout = QHBoxLayout(root)
        main_layout.setContentsMargins(0, 0, 0, 0)

        sidebar = QFrame()
        sidebar.setObjectName("Sidebar")
        sidebar.setFixedWidth(210)
        sidebar_layout = QVBoxLayout(sidebar)
        
        logo = QLabel("♻  ECOMIND AI")
        logo.setStyleSheet("font-size: 20px; font-weight: bold; color: #72f5c7; padding: 15px;")
        sidebar_layout.addWidget(logo)
        
        self.nav_buttons = []
        pages = [
            ("⌂   Dashboard", 0), ("＋   Add Waste", 1), ("▥   Waste Analytics", 2),
            ("↻   Recycling Tracker", 3), ("▤   History", 4), ("📊   Bar Graph", 5),
            ("🖼   Image Gallery", 6), ("✦   Insight Report", 7), ("⚗   What-if Lab", 8), ("⚙   Settings", 9)
        ]
        for text, index in pages:
            button = QPushButton(text)
            button.setObjectName("NavButton")
            button.clicked.connect(lambda checked=False, i=index: self.change_page(i))
            sidebar_layout.addWidget(button)
            self.nav_buttons.append(button)
        
        sidebar_layout.addStretch()
        status = QLabel("● AI Analysis Active\nEcoMind Intelligence Engine")
        status.setStyleSheet("color: #5ce0b2; padding: 8px 15px; font-size: 11px;")
        sidebar_layout.addWidget(status)

        account_button = QPushButton()
        account_button.setObjectName("AccountButton")
        account_button.setCursor(Qt.CursorShape.PointingHandCursor)
        account_button.setMinimumHeight(58)
        account_button.setMaximumHeight(64)
        account_button.setText("◉   Tech Inventor                 ⋮")
        account_menu = QMenu(self)
        account_menu.setStyleSheet("""
            QMenu { background: #0b1b2c; color: #eaf7ff; border: 1px solid #24527a; border-radius: 8px; padding: 6px; }
            QMenu::item { padding: 9px 28px 9px 12px; border-radius: 5px; }
            QMenu::item:selected { background: #173a5d; }
        """)
        profile_action = account_menu.addAction("  Tech Inventor")
        profile_action.setEnabled(False)
        account_menu.addSeparator()
        account_menu.addAction("⚙  Account Settings", lambda: self.change_page(9))
        account_menu.addAction("  About EcoMind AI", self.show_about)
        account_menu.addSeparator()
        account_menu.addAction("✕  Close Menu")
        account_button.clicked.connect(lambda: account_menu.exec(account_button.mapToGlobal(account_button.rect().topLeft())))
        sidebar_layout.addWidget(account_button)

        right = QWidget()
        right_layout = QVBoxLayout(right)
        right_layout.setContentsMargins(0, 0, 0, 0)

        topbar = QFrame()
        topbar.setObjectName("TopBar")
        top_layout = QHBoxLayout(topbar)
        title_box = QVBoxLayout()
        title = QLabel("EcoMind AI")
        title.setObjectName("Title")
        subtitle = QLabel("Intelligent Waste Management System")
        subtitle.setObjectName("Subtitle")
        title_box.addWidget(title)
        title_box.addWidget(subtitle)
        top_layout.addLayout(title_box)
        top_layout.addStretch()

        clock_box = QFrame()
        clock_box.setFixedSize(136, 48)
        clock_box.setStyleSheet("QFrame { background: qlineargradient(x1:0,y1:0,x2:1,y2:1, stop:0 #0B1A2D, stop:0.55 #101C35, stop:1 #17172F); border: 1px solid #304C72; border-radius: 10px; }")
        clock_layout = QHBoxLayout(clock_box)
        clock_layout.setContentsMargins(8, 5, 9, 5)
        clock_layout.setSpacing(7)
        clock_icon = QLabel("◷")
        clock_icon.setFixedSize(28, 28)
        clock_icon.setAlignment(Qt.AlignmentFlag.AlignCenter)
        clock_icon.setStyleSheet("background:#15294A;border:1px solid #647DFF;border-radius:14px;color:#78E7FF;font-size:15px;font-weight:bold;")
        clock_layout.addWidget(clock_icon)
        clock_text = QVBoxLayout()
        clock_text.setSpacing(0)
        self.clock_time_label = QLabel()
        self.clock_time_label.setStyleSheet("font-size:11px;font-weight:700;color:#F4F8FF;background:transparent;border:none;padding:0;")
        self.clock_date_label = QLabel()
        self.clock_date_label.setStyleSheet("font-size:8px;color:#8FA6C0;background:transparent;border:none;padding:0;")
        clock_text.addWidget(self.clock_time_label)
        clock_text.addWidget(self.clock_date_label)
        clock_layout.addLayout(clock_text)
        self.clock_widget = clock_box

        refresh_button = QPushButton("⟳ Refresh")
        export_button = QPushButton(" Export")
        settings_button = QPushButton("")
        about_button = QPushButton("ⓘ")
        refresh_button.clicked.connect(self.refresh_all)
        export_button.clicked.connect(self.export_docx)
        settings_button.clicked.connect(lambda: self.change_page(9))
        about_button.clicked.connect(self.show_about)
        
        top_layout.addWidget(self.clock_widget)
        top_layout.addWidget(refresh_button)
        top_layout.addWidget(export_button)
        top_layout.addWidget(settings_button)
        top_layout.addWidget(about_button)
        right_layout.addWidget(topbar)

        self.stack = QStackedWidget()
        self.dashboard_page = self.create_dashboard()
        self.add_page = self.create_add_waste()
        self.analytics_page = self.create_analytics()
        self.recycle_page = self.create_recycling()
        self.history_page = self.create_history()
        self.graph_page = self.create_graph_page()
        self.image_gallery_page = self.create_image_gallery_page()
        self.insight_report_page = self.create_insight_report()
        self.whatif_page = self.create_whatif_lab()
        self.settings_page = self.create_settings()
        
        self.stack.addWidget(self.dashboard_page)
        self.stack.addWidget(self.add_page)
        self.stack.addWidget(self.analytics_page)
        self.stack.addWidget(self.recycle_page)
        self.stack.addWidget(self.history_page)
        self.stack.addWidget(self.graph_page)
        self.stack.addWidget(self.image_gallery_page)
        self.stack.addWidget(self.insight_report_page)
        self.stack.addWidget(self.whatif_page)
        self.stack.addWidget(self.settings_page)
        
        right_layout.addWidget(self.stack)
        main_layout.addWidget(sidebar)
        main_layout.addWidget(right)
        self.setCentralWidget(root)
        self.change_page(0)

    def card(self, title, value):
        frame = QFrame()
        frame.setObjectName("Card")
        layout = QVBoxLayout(frame)
        title_label = QLabel(title)
        title_label.setObjectName("CardTitle")
        value_label = QLabel(value)
        value_label.setObjectName("CardValue")
        layout.addWidget(title_label)
        layout.addWidget(value_label)
        return frame, value_label

    def create_dashboard(self):
        page = QWidget()
        page.setStyleSheet("background:#020A16;")
        layout = QVBoxLayout(page)
        layout.setContentsMargins(10, 8, 10, 8)
        layout.setSpacing(8)

        metric_row = QHBoxLayout()
        metric_row.setSpacing(7)
        metric_data = [
            ("🗑", "TOTAL WASTE", "0.00 kg", "INPUT", "#25D9F5"),
            ("", "RECOVERED", "0.00 kg", "RECYCLE", "#35E6A1"),
            ("%", "RECOVERY RATE", "0.0%", "EFFICIENCY", "#5E91FF"),
            ("🍃", "TOP STREAM", "No data", "DOMINANT", "#FFC533"),
            ("", "ECO SCORE", "0.0", "INDEX", "#B477FF"),
            ("✦", "AI STATUS", "READY", "LOCAL ENGINE", "#FF4FA3"),
        ]
        refs = []
        for icon, label, value, sub, accent in metric_data:
            card = QFrame()
            card.setFixedHeight(56)
            card.setStyleSheet(f"background:qlineargradient(x1:0,y1:0,x2:1,y2:1,stop:0 #07182A,stop:1 #081224);border:1px solid #163B5A;border-top:2px solid {accent};border-radius:10px;")
            row = QHBoxLayout(card)
            row.setContentsMargins(7, 3, 7, 3)
            row.setSpacing(6)
            icon_box = QLabel(icon)
            icon_box.setFixedSize(28, 28)
            icon_box.setAlignment(Qt.AlignmentFlag.AlignCenter)
            icon_box.setStyleSheet(f"font-size:13px;font-weight:900;color:{accent};background:{accent}18;border:1px solid {accent}66;border-radius:15px;")
            row.addWidget(icon_box)
            text_col = QVBoxLayout()
            text_col.setSpacing(0)
            a = QLabel(label)
            a.setStyleSheet(f"font-size:7px;font-weight:900;letter-spacing:0.7px;color:{accent};background:transparent;border:none;")
            v = QLabel(value)
            v.setStyleSheet("font-size:15px;font-weight:900;color:#F4FAFF;background:transparent;border:none;")
            d = QLabel(sub)
            d.setStyleSheet("font-size:6px;color:#718AA4;background:transparent;border:none;")
            text_col.addWidget(a)
            text_col.addWidget(v)
            text_col.addWidget(d)
            row.addLayout(text_col, 1)
            metric_row.addWidget(card, 1)
            refs.append(v)
        
        self.total_value = refs[0]
        self.recycled_value = refs[1]
        self.rate_value = refs[2]
        self.highest_value = refs[3]
        self.ai_score_value = refs[4]
        self.ai_status_value = refs[5]
        self.reduction_value = QLabel("0.0")
        self.reduction_value.hide()
        layout.addLayout(metric_row, 0)

        flow_card = QFrame()
        flow_card.setStyleSheet("background:#03111F;border:1px solid #164C6B;border-radius:12px;")
        flow_layout = QVBoxLayout(flow_card)
        flow_layout.setContentsMargins(10, 8, 10, 8)
        flow_layout.setSpacing(3)
        flow_card.setMinimumHeight(420)
        flow_head = QHBoxLayout()
        flow_title = QLabel("🌿  ENVIRONMENTAL MATERIAL FLOW")
        flow_title.setStyleSheet("font-size:11px;font-weight:900;letter-spacing:1px;color:#EAF7FF;background:transparent;border:none;")
        flow_head.addWidget(flow_title)
        flow_head.addStretch()
        flow_note = QLabel("LIVE MATERIAL FLOW  /  LOCAL INTELLIGENCE")
        flow_note.setStyleSheet("font-size:7px;color:#6E8AA4;background:transparent;border:none;")
        flow_head.addWidget(flow_note)
        online = QLabel("● ONLINE")
        online.setStyleSheet("font-size:7px;font-weight:900;color:#35E6A1;background:#06261F;border:1px solid #1B6957;border-radius:6px;padding:4px 8px;")
        flow_head.addSpacing(8)
        flow_head.addWidget(online)
        flow_layout.addLayout(flow_head)
        
        self.flow_map = EcoFlowMap()
        self.flow_map.setMinimumHeight(375)
        self.flow_map.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)
        flow_layout.addWidget(self.flow_map, 1)
        layout.addWidget(flow_card, 1)

        analytics_row = QHBoxLayout()
        analytics_row.setSpacing(9)
        distribution = QFrame()
        distribution.setMinimumHeight(120)
        distribution.setMaximumHeight(126)
        distribution.setStyleSheet("background:#061423;border:1px solid #18445D;border-radius:11px;")
        dl = QVBoxLayout(distribution)
        dl.setContentsMargins(10, 7, 10, 7)
        dl.setSpacing(2)
        dt = QLabel("♻  WASTE STREAM DISTRIBUTION")
        dt.setStyleSheet("font-size:8px;font-weight:900;letter-spacing:1px;color:#58E8D0;background:transparent;border:none;")
        dl.addWidget(dt)
        self.dashboard_donut_figure = Figure(figsize=(3.2, 1.25), facecolor="#061423")
        self.dashboard_donut_canvas = FigureCanvas(self.dashboard_donut_figure)
        self.dashboard_donut_canvas.setMinimumHeight(82)
        self.dashboard_donut_canvas.setMaximumHeight(94)
        dl.addWidget(self.dashboard_donut_canvas)
        analytics_row.addWidget(distribution, 4)

        trend = QFrame()
        trend.setMinimumHeight(120)
        trend.setMaximumHeight(126)
        trend.setStyleSheet("background:#061423;border:1px solid #18445D;border-radius:11px;")
        tl = QVBoxLayout(trend)
        tl.setContentsMargins(9, 7, 9, 7)
        tl.setSpacing(2)
        tt = QLabel("▥  WASTE TREND MONITOR")
        tt.setStyleSheet("font-size:8px;font-weight:900;letter-spacing:1px;color:#62BFFF;background:transparent;border:none;")
        tl.addWidget(tt)
        self.dashboard_chart_figure = Figure(figsize=(4.0, 1.15), facecolor="#061423")
        self.dashboard_chart_canvas = FigureCanvas(self.dashboard_chart_figure)
        self.dashboard_chart_canvas.setMinimumHeight(82)
        self.dashboard_chart_canvas.setMaximumHeight(94)
        tl.addWidget(self.dashboard_chart_canvas)
        analytics_row.addWidget(trend, 5)

        impact = QFrame()
        impact.setMinimumHeight(120)
        impact.setMaximumHeight(126)
        impact.setStyleSheet("background:qlineargradient(x1:0,y1:0,x2:1,y2:1,stop:0 #071B28,stop:1 #12102A);border:1px solid #315277;border-radius:11px;")
        xl = QVBoxLayout(impact)
        xl.setContentsMargins(11, 8, 11, 8)
        xl.setSpacing(4)
        xt = QLabel("🍃  ECO IMPACT STATUS")
        xt.setStyleSheet("font-size:8px;font-weight:900;letter-spacing:1px;color:#5BE9C1;background:transparent;border:none;")
        xl.addWidget(xt)
        self.dashboard_distribution = QLabel("No data recorded yet")
        self.dashboard_distribution.setWordWrap(True)
        self.dashboard_distribution.setStyleSheet("font-size:9px;font-weight:800;color:#F0F7FF;background:transparent;border:none;")
        xl.addWidget(self.dashboard_distribution)
        xi = QLabel("Use the collected data to identify high-waste streams and improve recovery.")
        xi.setWordWrap(True)
        xi.setStyleSheet("font-size:7px;color:#7892AA;background:transparent;border:none;")
        xl.addWidget(xi)
        xl.addStretch()
        analytics_row.addWidget(impact, 2)
        layout.addLayout(analytics_row, 0)

        dock = QFrame()
        dock.setStyleSheet("background:#04111E;border:1px solid #16415D;border-radius:11px;")
        dock_layout = QHBoxLayout(dock)
        dock_layout.setContentsMargins(8, 4, 8, 4)
        dock_layout.setSpacing(7)
        dock_title = QLabel("QUICK ACTIONS")
        dock_title.setFixedWidth(120)
        dock_title.setStyleSheet("font-size:8px;font-weight:900;letter-spacing:1px;color:#91A8BD;background:transparent;border:none;")
        dock_layout.addWidget(dock_title)
        actions = [
            ("＋  Add Waste", lambda: self.change_page(1), "#35E6A1"),
            ("  Camera", lambda: self.change_page(1), "#25D9F5"),
            ("  Image Gallery", lambda: self.change_page(6), "#62BFFF"),
            ("▥  Analytics", lambda: self.change_page(2), "#FFC533"),
            ("♻  Recycling", lambda: self.change_page(3), "#35E6A1"),
        ]
        for text, command, accent in actions:
            btn = QPushButton(text)
            btn.setMinimumHeight(34)
            btn.setCursor(Qt.CursorShape.PointingHandCursor)
            btn.setStyleSheet(f"QPushButton{{background:#07182A;border:1px solid #1A4160;border-top:2px solid {accent};border-radius:8px;color:#DDEBFA;font-size:8px;font-weight:800;padding:6px 10px;}}QPushButton:hover{{background:{accent}20;border:1px solid {accent};color:white;}}")
            btn.clicked.connect(command)
            dock_layout.addWidget(btn, 1)
        layout.addWidget(dock, 0)
        return page

    def create_add_waste(self):
        page = QWidget()
        layout = QVBoxLayout(page)
        title = QLabel("Add Waste")
        title.setObjectName("SectionTitle")
        layout.addWidget(title)
        subtitle = QLabel("Enter waste manually or optionally use Computer Vision.")
        subtitle.setStyleSheet("color:#789d92;")
        layout.addWidget(subtitle)
        
        form = QGridLayout()
        form.addWidget(QLabel("House"), 0, 0)
        self.house_input = QComboBox()
        self.refresh_house_combo(self.house_input, include_all=False)
        form.addWidget(self.house_input, 0, 1)
        form.addWidget(QLabel("Date"), 1, 0)
        self.date_input = QLineEdit(datetime.now().strftime("%Y-%m-%d"))
        form.addWidget(self.date_input, 1, 1)
        form.addWidget(QLabel("Category"), 2, 0)
        self.category_input = QComboBox()
        self.category_input.addItems(CATEGORIES)
        form.addWidget(self.category_input, 2, 1)
        form.addWidget(QLabel("Weight (kg)"), 3, 0)
        self.weight_input = QDoubleSpinBox()
        self.weight_input.setRange(0.01, 10000)
        self.weight_input.setDecimals(2)
        form.addWidget(self.weight_input, 3, 1)
        form.addWidget(QLabel("Recycled"), 4, 0)
        self.recycled_input = QCheckBox("This waste was recycled")
        form.addWidget(self.recycled_input, 4, 1)
        form.addWidget(QLabel("Description"), 5, 0)
        self.description_input = QTextEdit()
        self.description_input.setMaximumHeight(100)
        form.addWidget(self.description_input, 5, 1)
        layout.addLayout(form)

        cv_box = QFrame()
        cv_box.setObjectName("Card")
        cv_layout = QHBoxLayout(cv_box)
        upload = QPushButton("🖼 Upload Waste Image")
        camera = QPushButton("📷 Open Camera")
        self.image_label = QLabel("No image selected")
        upload.clicked.connect(self.upload_image)
        camera.clicked.connect(self.capture_camera)
        cv_layout.addWidget(upload)
        cv_layout.addWidget(camera)
        cv_layout.addWidget(self.image_label)
        layout.addWidget(cv_box)
        
        self.ai_result_label = QLabel("AI Vision: waiting for image...")
        self.ai_result_label.setStyleSheet("color:#72f5c7;")
        layout.addWidget(self.ai_result_label)
        
        save_button = QPushButton("✓ Save Waste Record")
        save_button.setMinimumHeight(45)
        save_button.clicked.connect(self.save_waste)
        layout.addWidget(save_button)
        layout.addStretch()
        return page

    def upload_image(self):
        path, _ = QFileDialog.getOpenFileName(self, "Select Waste Image", "", "Images (*.png *.jpg *.jpeg)")
        if not path: return
        self.selected_image = path
        self.display_image(path)
        self.run_vision(path)

    def capture_camera(self):
        path = self.vision.capture_image()
        if not path: return
        self.selected_image = path
        self.display_image(path)
        self.run_vision(path)

    def display_image(self, path):
        pixmap = QPixmap(path)
        if pixmap.isNull(): return
        pixmap = pixmap.scaled(180, 120, Qt.AspectRatioMode.KeepAspectRatio, Qt.TransformationMode.SmoothTransformation)
        self.image_label.setPixmap(pixmap)

    def run_vision(self, image_path):
        self.ai_category = ""
        self.ai_confidence = 0.0
        try:
            category, confidence, message = self.ml_engine.classify_image(image_path)
            if category:
                self.ai_category = category
                self.ai_confidence = confidence
                self.ai_result_label.setText(f"AI Vision: {category}  •  {confidence:.1f}% confidence\n{message}")
            else:
                self.ai_result_label.setText(f"AI Vision: model is still learning.\n{message}\nSave more correctly labeled waste images to train the classifier.")
        except Exception as error:
            self.ai_result_label.setText(f"AI Vision unavailable: {error}")

    def save_waste(self):
        date = self.date_input.text().strip()
        category = self.category_input.currentText()
        weight = self.weight_input.value()
        recycled = self.recycled_input.isChecked()
        description = self.description_input.toPlainText().strip()
        house_id = self.house_input.currentData()
        if house_id is None: house_id = 1
        
        if not date:
            QMessageBox.warning(self, "Invalid Input", "Please enter a date.")
            return
        if weight <= 0:
            QMessageBox.warning(self, "Invalid Input", "Weight must be greater than zero.")
            return
        
        house_name = self.house_input.currentText()
        self.database.add_record(date, category, weight, recycled, description, self.selected_image, self.ai_category, self.ai_confidence, house_id)
        QMessageBox.information(self, "Saved", f"Waste record saved successfully for {house_name}.")
        
        self.weight_input.setValue(0)
        self.description_input.clear()
        self.recycled_input.setChecked(False)
        self.selected_image = ""
        self.ai_category = ""
        self.ai_confidence = 0
        self.image_label.setText("No image selected")
        self.load_dashboard()
        self.load_history()
        self.load_analytics()
        if hasattr(self, "load_insight_report"): self.load_insight_report()
        if hasattr(self, "run_whatif"): self.run_whatif()

    def create_analytics(self):
        page = QWidget()
        layout = QVBoxLayout(page)
        title = QLabel("Waste Analytics")
        title.setObjectName("SectionTitle")
        layout.addWidget(title)
        self.analytics_text = QTextEdit()
        self.analytics_text.setReadOnly(True)
        layout.addWidget(self.analytics_text)
        graph_button = QPushButton("📊  View Bar & Line Graphs")
        graph_button.setObjectName("PrimaryButton")
        graph_button.setMinimumHeight(44)
        graph_button.clicked.connect(self.open_graphs)
        layout.addWidget(graph_button)
        graph_hint = QLabel("Bar graph: waste by category   •   Line graph: waste generation over time")
        graph_hint.setStyleSheet("color:#86a99f; font-size:12px; padding:4px;")
        layout.addWidget(graph_hint)
        return page

    def load_analytics(self):
        data = self.analytics.category_breakdown()
        text = "CATEGORY BREAKDOWN\n\n"
        if not data:
            text += "No records available."
        else:
            for category, weight in sorted(data.items(), key=lambda x: x[1], reverse=True):
                text += f"{category:<15} {weight:.2f} kg\n"
            text += f"\nTotal Waste: {self.analytics.total_waste():.2f} kg\n"
            text += f"Recycled Waste: {self.analytics.recycled_waste():.2f} kg\n"
            text += f"Recycling Rate: {self.analytics.recycling_rate():.1f}%\n"
        self.analytics_text.setText(text)

    def create_recycling(self):
        page = QWidget()
        layout = QVBoxLayout(page)
        title = QLabel("Recycling Tracker")
        title.setObjectName("SectionTitle")
        layout.addWidget(title)
        self.recycling_progress = QProgressBar()
        layout.addWidget(self.recycling_progress)
        self.recycling_label = QLabel()
        layout.addWidget(self.recycling_label)
        layout.addStretch()
        return page

    def load_recycling(self):
        rate = self.analytics.recycling_rate()
        self.recycling_progress.setValue(int(rate))
        self.recycling_label.setText(f"Current recycling rate: {rate:.1f}%")

    def create_history(self):
        page = QWidget()
        layout = QVBoxLayout(page)
        title = QLabel("Waste History")
        title.setObjectName("SectionTitle")
        layout.addWidget(title)
        filter_row = QHBoxLayout()
        filter_row.addWidget(QLabel("Filter by House:"))
        self.history_house_filter = QComboBox()
        self.refresh_house_combo(self.history_house_filter, include_all=True)
        filter_row.addWidget(self.history_house_filter)
        filter_row.addStretch()
        layout.addLayout(filter_row)
        self.history_table = QTableWidget()
        self.history_table.setColumnCount(9)
        self.history_table.setHorizontalHeaderLabels(["ID", "House", "Date", "Category", "Weight", "Recycled", "Description", "Image", "AI Confidence"])
        self.history_table.horizontalHeader().setStretchLastSection(True)
        layout.addWidget(self.history_table)
        self.history_house_filter.currentIndexChanged.connect(self.load_history)
        return page

    def load_history(self):
        if not hasattr(self, "history_table"): return
        house_id = None
        if hasattr(self, "history_house_filter"): house_id = self.history_house_filter.currentData()
        records = self.database.get_records(house_id=house_id)
        self.history_table.setRowCount(len(records))
        for row_index, row in enumerate(records):
            values = [row[0], row[10] if len(row) > 10 else "—", row[1], row[2], f"{row[3]:.2f} kg", "Yes" if row[4] else "No", row[5] or "", "Yes" if row[6] else "No", f"{row[8]:.1f}%" if row[8] else "—"]
            for col, value in enumerate(values):
                self.history_table.setItem(row_index, col, QTableWidgetItem(str(value)))

    def create_graph_page(self):
        page = QWidget()
        layout = QVBoxLayout(page)
        title = QLabel("Waste Bar Graph")
        title.setObjectName("SectionTitle")
        layout.addWidget(title)
        subtitle = QLabel("Compare total waste generated by category.")
        subtitle.setStyleSheet("color:#86a99f; font-size:13px; padding-bottom:8px;")
        layout.addWidget(subtitle)
        self.graph_figure = Figure(figsize=(8, 5), facecolor="#101716")
        self.graph_canvas = FigureCanvas(self.graph_figure)
        layout.addWidget(self.graph_canvas)
        refresh = QPushButton("↻  Refresh Bar Graph")
        refresh.setObjectName("PrimaryButton")
        refresh.setMinimumHeight(42)
        refresh.clicked.connect(self.refresh_bar_graph)
        layout.addWidget(refresh)
        self.refresh_bar_graph()
        return page

    def refresh_bar_graph(self):
        self.graph_figure.clear()
        ax = self.graph_figure.add_subplot(111)
        ax.set_facecolor("#101716")
        try:
            cur = self.database.connection.cursor()
            cur.execute("SELECT category, COALESCE(SUM(weight), 0) FROM waste_records GROUP BY category ORDER BY SUM(weight) DESC")
            rows = cur.fetchall()
        except Exception as e:
            ax.text(0.5, 0.5, f"Could not load graph: {e}", ha="center", va="center", color="white")
            self.graph_canvas.draw()
            return
        if not rows:
            ax.text(0.5, 0.5, "No waste data available yet.", ha="center", va="center", fontsize=14, color="white")
            ax.set_axis_off()
        else:
            categories = [str(row[0]) for row in rows]
            weights = [float(row[1]) for row in rows]
            ax.bar(categories, weights)
            ax.set_title("Waste Generated by Category", fontsize=16, pad=15, color="white")
            ax.set_xlabel("Waste Category", color="white")
            ax.set_ylabel("Waste (kg)", color="white")
            ax.tick_params(axis="x", rotation=25, colors="white")
            ax.tick_params(axis="y", colors="white")
            ax.grid(axis="y", alpha=0.2)
            for spine in ax.spines.values(): spine.set_color("#3b5f55")
            self.graph_figure.tight_layout()
        self.graph_canvas.draw()

    def create_image_gallery_page(self):
        page = QWidget()
        layout = QVBoxLayout(page)
        title = QLabel("Waste Image Gallery")
        title.setObjectName("SectionTitle")
        layout.addWidget(title)
        subtitle = QLabel("View images saved with your waste records. Filter by house or view all.")
        subtitle.setStyleSheet("color:#86a99f; font-size:13px; padding-bottom:8px;")
        layout.addWidget(subtitle)
        filter_row = QHBoxLayout()
        filter_row.addWidget(QLabel("Filter by House:"))
        self.gallery_house_filter = QComboBox()
        self.refresh_house_combo(self.gallery_house_filter, include_all=True)
        filter_row.addWidget(self.gallery_house_filter)
        filter_row.addStretch()
        layout.addLayout(filter_row)
        self.image_scroll = QScrollArea()
        self.image_scroll.setWidgetResizable(True)
        self.image_scroll.setFrameShape(QFrame.Shape.NoFrame)
        self.image_gallery_container = QWidget()
        self.image_gallery_grid = QGridLayout(self.image_gallery_container)
        self.image_gallery_grid.setSpacing(16)
        self.image_scroll.setWidget(self.image_gallery_container)
        layout.addWidget(self.image_scroll)
        refresh = QPushButton("↻  Refresh Image Gallery")
        refresh.setObjectName("PrimaryButton")
        refresh.setMinimumHeight(42)
        refresh.clicked.connect(self.load_image_gallery)
        layout.addWidget(refresh)
        self.gallery_house_filter.currentIndexChanged.connect(self.load_image_gallery)
        self.load_image_gallery()
        return page

    def load_image_gallery(self):
        while self.image_gallery_grid.count():
            item = self.image_gallery_grid.takeAt(0)
            widget = item.widget()
            if widget is not None: widget.deleteLater()
        house_id = None
        if hasattr(self, "gallery_house_filter"): house_id = self.gallery_house_filter.currentData()
        records = self.database.get_records(house_id=house_id)
        image_records = [row for row in records if row[6] and os.path.exists(row[6])]
        if not image_records:
            empty = QLabel("No saved waste images found.\n\nAdd a waste record with Upload Waste Image or Open Camera, then save it.")
            empty.setAlignment(Qt.AlignmentFlag.AlignCenter)
            empty.setStyleSheet("color:#86a99f; font-size:15px; padding:40px;")
            self.image_gallery_grid.addWidget(empty, 0, 0, 1, 2)
            return
        for i, row in enumerate(image_records):
            record_id, date, category, weight, recycled, image_path = row[0], row[1], row[2], row[3], row[4], row[6]
            house_name = row[10] if len(row) > 10 else "—"
            card = QFrame()
            card.setObjectName("Card")
            card_layout = QVBoxLayout(card)
            image = QLabel()
            image.setAlignment(Qt.AlignmentFlag.AlignCenter)
            pixmap = QPixmap(image_path)
            if not pixmap.isNull():
                image.setPixmap(pixmap.scaled(250, 170, Qt.AspectRatioMode.KeepAspectRatio, Qt.TransformationMode.SmoothTransformation))
            else:
                image.setText("Image could not be loaded")
            image.setMinimumSize(250, 170)
            card_layout.addWidget(image)
            info = QLabel(f"🏠 {house_name}\nRecord #{record_id}  •  {date}\n{category}  •  {float(weight):.2f} kg\n{'Recycled' if recycled else 'Not recycled'}")
            info.setStyleSheet("color:#d7eee7; font-size:12px; padding:4px;")
            card_layout.addWidget(info)
            view_button = QPushButton("  View Image")
            view_button.clicked.connect(lambda checked=False, path=image_path, rid=record_id: self.view_saved_image(path, rid))
            card_layout.addWidget(view_button)
            column = i % 2
            row_number = i // 2
            self.image_gallery_grid.addWidget(card, row_number, column)
        self.image_gallery_grid.setColumnStretch(0, 1)
        self.image_gallery_grid.setColumnStretch(1, 1)

    def view_saved_image(self, image_path, record_id):
        dialog = QDialog(self)
        dialog.setWindowTitle(f"EcoMind AI — Waste Image #{record_id}")
        dialog.resize(900, 650)
        layout = QVBoxLayout(dialog)
        image_label = QLabel()
        image_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        pixmap = QPixmap(image_path)
        if pixmap.isNull():
            image_label.setText("Could not open this image.")
            image_label.setStyleSheet("color:white; font-size:16px;")
        else:
            image_label.setPixmap(pixmap.scaled(820, 540, Qt.AspectRatioMode.KeepAspectRatio, Qt.TransformationMode.SmoothTransformation))
        layout.addWidget(image_label)
        close_button = QPushButton("Close")
        close_button.clicked.connect(dialog.accept)
        layout.addWidget(close_button)
        dialog.exec()

    def create_insight_report(self):
        page = QWidget()
        layout = QVBoxLayout(page)
        title = QLabel("✦ Insight Report")
        title.setObjectName("SectionTitle")
        layout.addWidget(title)
        subtitle = QLabel("Summary of waste patterns, high/low streams, and recommended actions.")
        subtitle.setStyleSheet("color:#86a99f; font-size:13px; padding-bottom:8px;")
        layout.addWidget(subtitle)
        self.insight_report_text = QTextEdit()
        self.insight_report_text.setReadOnly(True)
        self.insight_report_text.setStyleSheet("background:#0a1620; color:#d7eee7; border:1px solid #1C3A59; border-radius:8px; padding:12px; font-size:13px;")
        layout.addWidget(self.insight_report_text)
        refresh_btn = QPushButton("↻  Refresh Insight Report")
        refresh_btn.setObjectName("PrimaryButton")
        refresh_btn.setMinimumHeight(42)
        refresh_btn.clicked.connect(self.load_insight_report)
        layout.addWidget(refresh_btn)
        self.load_insight_report()
        return page

    def load_insight_report(self):
        if not hasattr(self, "insight_report_text"): return
        records = self.database.get_records()
        if not records:
            self.insight_report_text.setText("ECOMIND INSIGHT REPORT\n────────────────────────────\n\nNo waste records yet.\nAdd waste entries to generate insights on high and low waste streams.")
            return
        total = self.analytics.total_waste()
        recycled = self.analytics.recycled_waste()
        rate = self.analytics.recycling_rate()
        breakdown = self.analytics.category_breakdown()
        highest = self.analytics.highest_category()
        sorted_cats = sorted(breakdown.items(), key=lambda x: x[1], reverse=True)
        high_cats = [(c, w) for c, w in sorted_cats if w >= HIGH_WASTE_THRESHOLD_KG]
        low_cats = [(c, w) for c, w in sorted_cats if w < HIGH_WASTE_THRESHOLD_KG]
        house_totals = {}
        for row in records:
            hname = row[10] if len(row) > 10 else "Unknown"
            house_totals[hname] = house_totals.get(hname, 0.0) + float(row[3] or 0)
        
        lines = ["ECOMIND INSIGHT REPORT", "────────────────────────────", f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}", "", "OVERVIEW", f"  Total waste recorded : {total:.2f} kg", f"  Recycled             : {recycled:.2f} kg", f"  Recycling rate       : {rate:.1f}%", f"  Records analyzed     : {len(records)}", f"  Dominant stream      : {highest}", "", "HIGH WASTE STREAMS  (≥ 10 kg)"]
        if high_cats:
            for cat, w in high_cats: lines.append(f"  ⚠  {cat}: {w:.2f} kg")
        else:
            lines.append("  ✓  No category currently exceeds 10 kg.")
        lines.append("")
        lines.append("LOW / MODERATE STREAMS  (< 10 kg)")
        if low_cats:
            for cat, w in low_cats: lines.append(f"  •  {cat}: {w:.2f} kg")
        else:
            lines.append("  (none)")
        lines.append("")
        lines.append("BY HOUSE")
        for hname, w in sorted(house_totals.items(), key=lambda x: -x[1]):
            flag = " ⚠ HIGH" if w >= HIGH_WASTE_THRESHOLD_KG else ""
            lines.append(f"  🏠 {hname}: {w:.2f} kg{flag}")
        lines.append("")
        lines.append("RECOMMENDED ACTIONS")
        if high_cats:
            for cat, w in high_cats:
                tips = HIGH_WASTE_SUGGESTIONS.get(cat, HIGH_WASTE_SUGGESTIONS["General"])
                lines.append(f"  [{cat} — {w:.2f} kg]")
                for tip in tips[:2]: lines.append(f"    → {tip}")
        else:
            for tip in HIGH_WASTE_SUGGESTIONS["General"]: lines.append(f"  → {tip}")
        if rate < 50:
            lines.append("  → Recycling rate is under 50%. Improve segregation at source.")
        else:
            lines.append("  → Good recycling rate. Keep tracking daily to maintain progress.")
        self.insight_report_text.setText("\n".join(lines))

    def create_whatif_lab(self):
        page = QWidget()
        layout = QVBoxLayout(page)
        title = QLabel("⚗ What-if Lab")
        title.setObjectName("SectionTitle")
        layout.addWidget(title)
        subtitle = QLabel("Explore high vs low waste scenarios. Adjust the reduction slider to see projected impact.")
        subtitle.setStyleSheet("color:#86a99f; font-size:13px; padding-bottom:8px;")
        layout.addWidget(subtitle)
        controls = QFrame()
        controls.setObjectName("Card")
        controls_layout = QVBoxLayout(controls)
        row1 = QHBoxLayout()
        row1.addWidget(QLabel("Scenario focus:"))
        self.whatif_category = QComboBox()
        self.whatif_category.addItem("All categories", "ALL")
        for c in CATEGORIES: self.whatif_category.addItem(c, c)
        self.whatif_category.currentIndexChanged.connect(self.run_whatif)
        row1.addWidget(self.whatif_category)
        row1.addStretch()
        controls_layout.addLayout(row1)
        row2 = QHBoxLayout()
        row2.addWidget(QLabel("Waste reduction %:"))
        self.whatif_slider = QSlider(Qt.Orientation.Horizontal)
        self.whatif_slider.setRange(0, 80)
        self.whatif_slider.setValue(30)
        self.whatif_slider.setTickPosition(QSlider.TickPosition.TicksBelow)
        self.whatif_slider.setTickInterval(10)
        self.whatif_slider.valueChanged.connect(self.run_whatif)
        row2.addWidget(self.whatif_slider)
        self.whatif_pct_label = QLabel("30%")
        self.whatif_pct_label.setFixedWidth(40)
        row2.addWidget(self.whatif_pct_label)
        controls_layout.addLayout(row2)
        layout.addWidget(controls)
        self.whatif_result = QTextEdit()
        self.whatif_result.setReadOnly(True)
        self.whatif_result.setStyleSheet("background:#0a1620; color:#d7eee7; border:1px solid #1C3A59; border-radius:8px; padding:12px; font-size:13px;")
        layout.addWidget(self.whatif_result)
        run_btn = QPushButton("▶  Run What-if Scenario")
        run_btn.setObjectName("PrimaryButton")
        run_btn.setMinimumHeight(42)
        run_btn.clicked.connect(self.run_whatif)
        layout.addWidget(run_btn)
        self.run_whatif()
        return page

    def run_whatif(self):
        if not hasattr(self, "whatif_result"): return
        pct = self.whatif_slider.value() if hasattr(self, "whatif_slider") else 30
        if hasattr(self, "whatif_pct_label"): self.whatif_pct_label.setText(f"{pct}%")
        focus = "ALL"
        if hasattr(self, "whatif_category"): focus = self.whatif_category.currentData() or "ALL"
        breakdown = self.analytics.category_breakdown()
        total = self.analytics.total_waste()
        if total <= 0:
            self.whatif_result.setText("WHAT-IF LAB\n────────────────────────────\n\nNo waste data yet.\nAdd records first, then simulate high and low waste scenarios here.")
            return
        factor = 1.0 - (pct / 100.0)
        lines = ["WHAT-IF LAB — SCENARIO ANALYSIS", "────────────────────────────", f"Reduction applied : {pct}%", f"Focus             : {focus if focus != 'ALL' else 'All categories'}", "", "CURRENT  (baseline)", f"  Total waste : {total:.2f} kg"]
        for cat in CATEGORIES:
            w = breakdown.get(cat, 0.0)
            marker = " ← HIGH" if w >= HIGH_WASTE_THRESHOLD_KG else ""
            lines.append(f"  {cat:10s}: {w:7.2f} kg{marker}")
        lines.append("")
        lines.append(f"LOW-WASTE SCENARIO  (−{pct}%)")
        projected_total = 0.0
        for cat in CATEGORIES:
            w = breakdown.get(cat, 0.0)
            new_w = w * factor if (focus == "ALL" or focus == cat) else w
            projected_total += new_w
            marker = " ← still high" if new_w >= HIGH_WASTE_THRESHOLD_KG else " ✓"
            lines.append(f"  {cat:10s}: {new_w:7.2f} kg{marker}")
        saved = total - projected_total
        lines.append("")
        lines.append(f"  Projected total : {projected_total:.2f} kg")
        lines.append(f"  Waste avoided   : {saved:.2f} kg")
        if total > 0: lines.append(f"  Relative cut    : {(saved / total) * 100:.1f}%")
        lines.append("")
        lines.append("HIGH-WASTE SCENARIO  (+25% growth if no action)")
        growth = 1.25
        high_total = 0.0
        for cat in CATEGORIES:
            w = breakdown.get(cat, 0.0)
            new_w = w * growth
            high_total += new_w
            marker = " ⚠" if new_w >= HIGH_WASTE_THRESHOLD_KG else ""
            lines.append(f"  {cat:10s}: {new_w:7.2f} kg{marker}")
        lines.append(f"  Projected total : {high_total:.2f} kg")
        lines.append(f"  Extra waste     : {high_total - total:.2f} kg")
        lines.append("")
        lines.append("TAKEAWAY")
        if saved > 0:
            lines.append(f"  Reducing by {pct}% could avoid {saved:.2f} kg of waste. Start with the highest streams shown above.")
        else:
            lines.append("  Move the slider to model a reduction and compare high vs low outcomes.")
        self.whatif_result.setText("\n".join(lines))

    def create_settings(self):
        page = QWidget()
        layout = QVBoxLayout(page)
        title = QLabel("Settings")
        title.setObjectName("SectionTitle")
        layout.addWidget(title)
        house_card = QFrame()
        house_card.setObjectName("Card")
        house_layout = QVBoxLayout(house_card)
        house_title = QLabel("🏠 Manage Houses")
        house_title.setStyleSheet("font-size:15px; font-weight:bold; color:#72f5c7;")
        house_layout.addWidget(house_title)
        self.houses_list_label = QLabel()
        self.houses_list_label.setStyleSheet("color:#d7eee7; padding:6px;")
        house_layout.addWidget(self.houses_list_label)
        add_row = QHBoxLayout()
        self.new_house_input = QLineEdit()
        self.new_house_input.setPlaceholderText("e.g. Mehta House")
        add_row.addWidget(self.new_house_input)
        add_house_btn = QPushButton("＋ Add House")
        add_house_btn.clicked.connect(self.add_new_house)
        add_row.addWidget(add_house_btn)
        house_layout.addLayout(add_row)
        layout.addWidget(house_card)
        self.refresh_houses_list_label()
        clear_button = QPushButton("Clear All Waste Data")
        clear_button.clicked.connect(self.clear_data)
        layout.addWidget(clear_button)
        info = QLabel("Database: Local SQLite\nAI Vision: OpenCV pipeline ready\nHouses: Multi-house waste tracking enabled")
        info.setStyleSheet("color:#7faaa0; padding:15px;")
        layout.addWidget(info)
        layout.addStretch()
        return page

    def refresh_house_combo(self, combo, include_all=True):
        current = combo.currentData() if combo.count() else None
        while combo.count(): combo.removeItem(0)
        if include_all: combo.addItem("All Houses", None)
        for house_id, name in self.database.get_houses(): combo.addItem(f"{name}  (ID-{house_id})", house_id)
        if current is not None:
            idx = combo.findData(current)
            if idx >= 0: combo.setCurrentIndex(idx)

    def refresh_houses_list_label(self):
        houses = self.database.get_houses()
        if not houses:
            self.houses_list_label.setText("No houses yet.")
            return
        lines = [f"• {name}  —  ID-{hid}" for hid, name in houses]
        self.houses_list_label.setText("\n".join(lines))

    def add_new_house(self):
        name = self.new_house_input.text().strip()
        ok, result = self.database.add_house(name)
        if not ok:
            QMessageBox.warning(self, "Add House", result)
            return
        self.new_house_input.clear()
        self.refresh_houses_list_label()
        if hasattr(self, "house_input"): self.refresh_house_combo(self.house_input, include_all=False)
        if hasattr(self, "history_house_filter"):
            self.history_house_filter.blockSignals(True)
            self.refresh_house_combo(self.history_house_filter, include_all=True)
            self.history_house_filter.blockSignals(False)
        if hasattr(self, "gallery_house_filter"):
            self.gallery_house_filter.blockSignals(True)
            self.refresh_house_combo(self.gallery_house_filter, include_all=True)
            self.gallery_house_filter.blockSignals(False)
        QMessageBox.information(self, "House Added", f"'{name}' has been added (ID-{result}).")

    def change_page(self, index):
        self.stack.setCurrentIndex(index)
        for i, button in enumerate(self.nav_buttons):
            if i == index: button.setObjectName("ActiveNav")
            else: button.setObjectName("NavButton")
            button.style().unpolish(button)
            button.style().polish(button)
        if index == 0: self.load_dashboard()
        elif index == 2: self.load_analytics()
        elif index == 3: self.load_recycling()
        elif index == 4: self.load_history()
        elif index == 6: self.load_image_gallery()

    def load_dashboard(self):
        total = self.analytics.total_waste()
        recycled = self.analytics.recycled_waste()
        rate = self.analytics.recycling_rate()
        highest = self.analytics.highest_category()
        score = self.analytics.sustainability_score()
        self.total_value.setText(f"{total:.2f} kg")
        self.recycled_value.setText(f"{recycled:.2f} kg")
        self.rate_value.setText(f"{rate:.1f}%")
        self.highest_value.setText(highest)
        self.ai_score_value.setText(f"{score:.1f}")
        self.reduction_value.setText(f"{score:.1f}")
        if hasattr(self, "flow_map"): self.flow_map.set_data(self.analytics.category_breakdown(), score)
        
        if hasattr(self, "dashboard_donut_figure"):
            try:
                self.dashboard_donut_figure.clear()
                ax = self.dashboard_donut_figure.add_subplot(111)
                ax.set_facecolor("#061423")
                breakdown = self.analytics.category_breakdown()
                labels = [cat for cat in CATEGORIES if float(breakdown.get(cat, 0) or 0) > 0]
                values = [float(breakdown.get(cat, 0) or 0) for cat in labels]
                if values and sum(values) > 0:
                    palette = {"Food": "#35E6A1", "Plastic": "#25D9F5", "Paper": "#FFC533", "Glass": "#5E91FF", "Metal": "#B477FF", "E-waste": "#FF4FA3", "Other": "#61E6D2"}
                    wedges, _ = ax.pie(values, startangle=90, counterclock=False, colors=[palette.get(c, "#55758E") for c in labels], wedgeprops={"width": 0.34, "edgecolor": "#061423", "linewidth": 2.2})
                    ax.legend(wedges, labels, loc="center left", bbox_to_anchor=(0.98, 0.5), frameon=False, labelcolor="#DDEBFA", fontsize=7.2, handlelength=0.9, handletextpad=0.5, borderaxespad=0.0)
                else:
                    ax.text(0.5, 0.54, "NO WASTE RECORDED YET", ha="center", va="center", color="#7895AE", fontsize=8.5, fontweight="bold", transform=ax.transAxes)
                    ax.text(0.5, 0.43, "Add your first waste entry to begin analysis", ha="center", va="center", color="#506B83", fontsize=6.5, transform=ax.transAxes)
                ax.set_axis_off()
                self.dashboard_donut_figure.tight_layout(pad=0.2)
                self.dashboard_donut_canvas.draw_idle()
            except Exception: pass
        
        if hasattr(self, "dashboard_chart_figure"):
            try:
                self.dashboard_chart_figure.clear()
                ax = self.dashboard_chart_figure.add_subplot(111)
                ax.set_facecolor("#071523")
                records = self.database.get_records()
                dates = {}
                for row in records: dates[row[1]] = dates.get(row[1], 0) + float(row[3] or 0)
                items = sorted(dates.items())[-7:]
                if items:
                    x = list(range(len(items)))
                    y = [v for _, v in items]
                    ax.plot(x, y, marker="o", linewidth=2, markersize=4, color="#31D8FF")
                    ax.fill_between(x, y, [0] * len(y), alpha=0.10, color="#31D8FF")
                    ax.set_xticks(x)
                    ax.set_xticklabels([d[-5:] for d, _ in items], color="#7E98B0", fontsize=6)
                    ax.tick_params(axis="y", colors="#7E98B0", labelsize=6)
                    ax.grid(axis="y", alpha=0.14, color="#58738A")
                    ax.spines[:].set_visible(False)
                else:
                    ax.text(0.5, 0.54, "NO HISTORICAL DATA YET", ha="center", va="center", color="#7895AE", fontsize=8, fontweight="bold", transform=ax.transAxes)
                    ax.text(0.5, 0.43, "Your daily waste trend will appear here", ha="center", va="center", color="#506B83", fontsize=6.5, transform=ax.transAxes)
                    ax.set_xticks([])
                    ax.set_yticks([])
                self.dashboard_chart_figure.tight_layout(pad=0.4)
                self.dashboard_chart_canvas.draw_idle()
            except Exception: pass
        
        breakdown = self.analytics.category_breakdown()
        if hasattr(self, "dashboard_distribution"):
            if breakdown:
                ranked = sorted(breakdown.items(), key=lambda item: item[1], reverse=True)
                self.dashboard_distribution.setText("   •   ".join(f"{cat}: {value:.1f} kg" for cat, value in ranked[:4]))
            else:
                self.dashboard_distribution.setText("No data recorded yet")
        
        if total == 0:
            self.ai_status_value.setText("READY")
        else:
            try:
                ml = self.ml_engine.analyze()
                active = ml["active_models"]
                self.ai_status_value.setText("ACTIVE" if active else "READY")
            except Exception:
                self.ai_status_value.setText("ERROR")

    def refresh_all(self):
        self.load_dashboard()
        self.load_analytics()
        self.load_recycling()
        self.load_history()
        if hasattr(self, "load_insight_report"): self.load_insight_report()
        if hasattr(self, "run_whatif"): self.run_whatif()

    def export_docx(self):
        records = self.database.get_records()
        if not records:
            QMessageBox.information(self, "Export", "There is no data to export.")
            return
        path, _ = QFileDialog.getSaveFileName(self, "Export Waste Report", "ecomind_waste_report.docx", "Word Document (*.docx)")
        if not path: return
        if not path.lower().endswith(".docx"): path += ".docx"
        try:
            doc = Document()
            title = doc.add_heading("EcoMind AI — Waste Report", level=0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            meta = doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}\nTotal records: {len(records)}")
            meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_heading("Summary", level=1)
            total = self.analytics.total_waste()
            recycled = self.analytics.recycled_waste()
            rate = self.analytics.recycling_rate()
            highest = self.analytics.highest_category()
            summary = doc.add_paragraph()
            summary.add_run(f"Total waste: {total:.2f} kg\n")
            summary.add_run(f"Recycled: {recycled:.2f} kg\n")
            summary.add_run(f"Recycling rate: {rate:.1f}%\n")
            summary.add_run(f"Dominant category: {highest}")
            doc.add_heading("Waste by Category", level=1)
            breakdown = self.analytics.category_breakdown()
            for cat, weight in sorted(breakdown.items(), key=lambda x: -x[1]): doc.add_paragraph(f"{cat}: {weight:.2f} kg", style="List Bullet")
            doc.add_heading("Waste by House", level=1)
            house_totals = {}
            for row in records:
                hname = row[10] if len(row) > 10 else "Unknown"
                house_totals[hname] = house_totals.get(hname, 0.0) + float(row[3] or 0)
            for hname, weight in sorted(house_totals.items(), key=lambda x: -x[1]): doc.add_paragraph(f"{hname}: {weight:.2f} kg", style="List Bullet")
            doc.add_heading("Detailed Records", level=1)
            table = doc.add_table(rows=1, cols=7)
            table.style = "Table Grid"
            headers = ["ID", "House", "Date", "Category", "Weight (kg)", "Recycled", "Description"]
            header_cells = table.rows[0].cells
            for i, h in enumerate(headers):
                header_cells[i].text = h
                for paragraph in header_cells[i].paragraphs:
                    for run in paragraph.runs: run.bold = True
            for row in records:
                cells = table.add_row().cells
                cells[0].text = str(row[0])
                cells[1].text = str(row[10] if len(row) > 10 else "—")
                cells[2].text = str(row[1])
                cells[3].text = str(row[2])
                cells[4].text = f"{float(row[3]):.2f}"
                cells[5].text = "Yes" if row[4] else "No"
                cells[6].text = str(row[5] or "")
            doc.add_paragraph("")
            footer = doc.add_paragraph("Report generated by EcoMind AI — Intelligent Waste Management System")
            footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.save(path)
            QMessageBox.information(self, "Export Complete", f"Waste report saved as Word document:\n{path}")
        except Exception as error:
            QMessageBox.critical(self, "Export Error", str(error))

    def clear_data(self):
        answer = QMessageBox.question(self, "Confirm", "Delete ALL waste records?", QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No)
        if answer != QMessageBox.StandardButton.Yes: return
        self.database.clear_all()
        self.refresh_all()
        QMessageBox.information(self, "Complete", "All waste data has been cleared.")

    def show_about(self):
        QMessageBox.information(self, "About EcoMind AI", "EcoMind AI\n\nIntelligent Waste Management System\n\nCore mission:\nUse AI and data intelligence to understand waste, detect patterns, predict trends, and help users reduce environmental impact.\n\nPython • PyQt6 • SQLite • OpenCV • scikit-learn • NumPy")

    def closeEvent(self, event):
        self.database.close()
        event.accept()

# ============================================================
# APPLICATION START
# ============================================================
if __name__ == "__main__":
    app = QApplication(sys.argv)
    app.setFont(QFont("Segoe UI", 10))
    window = EcoMindApp()
    window.show()
    sys.exit(app.exec())